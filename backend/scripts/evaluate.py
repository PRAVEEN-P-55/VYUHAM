"""Evaluation harness (spec 8).

Runs the live pipeline functions (not over HTTP) against the six planted
ground-truth files and prints a scorecard. This is the judge-facing proof
that the analytics actually recover what was planted -- and, for pattern
detection, that they do NOT fire on the 20 deliberate false-positive traps.

    python scripts/evaluate.py
"""
from __future__ import annotations

import json
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

import networkx as nx  # noqa: E402

from app.config import settings  # noqa: E402
from app.services import graph_store, mo_similarity, patterns  # noqa: E402
from app.services.identity_resolution import score_pair  # noqa: E402

DS = settings.dataset_dir


def _jsonl(name: str) -> list[dict]:
    return [json.loads(l) for l in (DS / name).open(encoding="utf-8") if l.strip()]


def _pct(x: float) -> str:
    return f"{100 * x:5.1f}%"


def _prf(tp: int, fp: int, fn: int) -> tuple[float, float]:
    precision = tp / (tp + fp) if (tp + fp) else 0.0
    recall = tp / (tp + fn) if (tp + fn) else 0.0
    return precision, recall


# ---------------------------------------------------------------------------
def eval_network() -> str:
    rows = _jsonl("network_ground_truth.jsonl")
    tp = fp = fn = tn = 0
    for r in rows:
        support = graph_store.global_edge_support(
            r["source_entity_id"], r["target_entity_id"], r["relationship_type"]
        )
        recovered = support >= r.get("minimum_supporting_sources", 1)
        expected = bool(r["expected_relationship"])
        if expected and recovered:
            tp += 1
        elif expected and not recovered:
            fn += 1
        elif not expected and recovered:
            fp += 1
        else:
            tn += 1
    p, rec = _prf(tp, fp, fn)
    n_true = sum(1 for r in rows if r["expected_relationship"])
    return (f"NETWORK RECOVERY     : precision {_pct(p)}  recall {_pct(rec)}   "
            f"({len(rows)} rows: {n_true} TRUE / {len(rows) - n_true} FALSE; "
            f"tp={tp} fp={fp} fn={fn} tn={tn})")


def eval_influencer() -> str:
    """Person-level influence: for each planted (entity, centrality_type, reason)
    row, check the entity lands in the top 3 of that centrality in its scenario.

    EIGENVECTOR / DEGREE are read off the tight scenario core (members + org +
    broker, phones/accounts folded into owners); BETWEENNESS needs bridge
    context so it is read off the neighbour-inclusive graph.
    """
    rows = _jsonl("influencer_ground_truth.jsonl")
    clusters = {c["cluster_id"]: c for c in _jsonl("scenario_cluster_summary.jsonl")}
    graph_store.ensure_loaded()

    def pick(cid: str, keys: tuple[str, ...]) -> list[str]:
        c = clusters.get(cid, {})
        out: list[str] = []
        for k in keys:
            v = c.get(k)
            out += (v if isinstance(v, list) else [v]) if v else []
        return out

    TIGHT = ("person_ids", "org_ids", "broker_phone_ids")
    WIDE = ("person_ids", "phone_ids", "account_ids", "org_ids", "broker_phone_ids")

    hits = 0
    per_type: dict[str, list[int]] = {}
    cache: dict[str, dict] = {}
    for r in rows:
        cid = r["scenario_id"]
        if cid not in cache:
            tight = graph_store.subgraph_from_entities(pick(cid, TIGHT), False, True)
            wide = graph_store.subgraph_from_entities(pick(cid, WIDE), True, True)
            d: dict = {}
            if tight.number_of_nodes():
                try:
                    d["EIGENVECTOR"] = nx.eigenvector_centrality_numpy(tight)
                except Exception:
                    d["EIGENVECTOR"] = nx.degree_centrality(tight)
                d["DEGREE"] = nx.degree_centrality(tight)
            if wide.number_of_nodes():
                d["BETWEENNESS"] = nx.betweenness_centrality(wide, weight="weight")
            cache[cid] = d
        scores = cache[cid].get(r["centrality_type"])
        bucket = per_type.setdefault(r["centrality_type"], [0, 0])
        bucket[1] += 1
        if not scores:
            continue
        ranking = [n for n, _ in sorted(scores.items(), key=lambda kv: kv[1], reverse=True)]
        if r["entity_id"] in ranking[:3]:
            hits += 1
            bucket[0] += 1
    detail = "  ".join(f"{k} {a}/{b}" for k, (a, b) in sorted(per_type.items()))
    return (f"INFLUENCER RANKING   : top-3 hit rate {_pct(hits / len(rows))}          "
            f"({len(rows)} rows across {len(set(r['influence_reason'] for r in rows))} reasons)\n"
            f"    {detail}")


def eval_mo() -> str:
    rows = _jsonl("mo_ground_truth.jsonl")
    incidents = {
        i["incident_id"]: i for i in _jsonl_or_db_incidents()
    }
    abs_err = 0.0
    n = 0
    correct_group = 0
    for r in rows:
        a, b = incidents.get(r["incident_a_id"]), incidents.get(r["incident_b_id"])
        if not a or not b:
            continue
        got = mo_similarity.compare(a, b)["similarity_pct"]
        abs_err += abs(got - r["expected_similarity"])
        n += 1
        if (got >= 0.5) == bool(r["same_pattern_group"]):
            correct_group += 1
    mae = abs_err / n if n else 0.0
    return (f"MO SIMILARITY        : mean abs error {mae:.3f} vs expected   "
            f"({n} rows; same-group agreement {_pct(correct_group / n) if n else 'n/a'})")


def _jsonl_or_db_incidents() -> list[dict]:
    from app.services.db import query
    return [dict(r) for r in query("SELECT * FROM incidents")]


def eval_patterns() -> str:
    rows = _jsonl("suspicious_pattern_ground_truth.jsonl")
    tp = fp = fn = tn = 0
    per_type: dict[str, list[int]] = {}
    for r in rows:
        fired = patterns.evaluate_row(
            r["pattern_type"], r.get("entity_ids_involved", []),
            r.get("supporting_record_ids", []),
        )
        expected = bool(r["expected_flag"])
        bucket = per_type.setdefault(r["pattern_type"], [0, 0, 0, 0])
        if expected and fired:
            tp += 1; bucket[0] += 1
        elif expected and not fired:
            fn += 1; bucket[2] += 1
        elif not expected and fired:
            fp += 1; bucket[1] += 1
        else:
            tn += 1; bucket[3] += 1
    p, rec = _prf(tp, fp, fn)
    n_true = sum(1 for r in rows if r["expected_flag"])
    lines = [f"PATTERN DETECTION    : precision {_pct(p)}  recall {_pct(rec)}    "
             f"({len(rows)} rows: {n_true} real / {len(rows) - n_true} traps)"]
    for pt, (a, b, c, d) in sorted(per_type.items()):
        lines.append(f"    - {pt:<24} tp={a} fp={b} fn={c} tn={d}")
    return "\n".join(lines)


def eval_gaps() -> str:
    from app.services.db import query
    gt = _jsonl("evidence_gaps.jsonl")
    planted_loc = [g for g in gt if g["missing_evidence_type"] == "MISSING_LOCATION_DATA"]
    served = {r["gap_id"] for r in query("SELECT gap_id FROM evidence_gaps")}
    recovered = sum(1 for g in gt if g["gap_id"] in served)
    loc_recovered = sum(1 for g in planted_loc if g["gap_id"] in served)
    return (f"EVIDENCE GAPS        : recall {_pct(recovered / len(gt))} on planted gaps  "
            f"({len(gt)} rows; {loc_recovered}/{len(planted_loc)} planted missing-location)")


def eval_identity() -> str:
    from openpyxl import load_workbook
    from app.services.db import query_one, row_to_dict

    wb = load_workbook(DS / "identity_ground_truth.xlsx", read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(it)]
    correct = total = skipped = 0
    for raw in it:
        row = dict(zip(header, raw))
        a = row_to_dict(query_one("SELECT * FROM people WHERE person_id = ?", (row["record_a_id"],)))
        b = row_to_dict(query_one("SELECT * FROM people WHERE person_id = ?", (row["record_b_id"],)))
        if not a or not b:
            skipped += 1
            continue
        pred = score_pair(a, b)["same_person"]
        truth = str(row["same_person"]).strip().upper() in ("TRUE", "1", "YES")
        total += 1
        correct += int(pred == truth)
    acc = correct / total if total else 0.0
    return (f"IDENTITY RESOLUTION  : accuracy {_pct(acc)}                 "
            f"({total} person-person rows scored, {skipped} mention rows skipped)")


def main() -> None:
    if not settings.db_path.exists():
        sys.exit("No DB. Run: python scripts/seed_from_datasets.py")
    graph_store.ensure_loaded()
    print("\n" + "=" * 72)
    print("  VYUHAM INTELLIGENCE  --  GROUND-TRUTH SCORECARD")
    print("=" * 72)
    for fn in (eval_network, eval_influencer, eval_mo, eval_patterns, eval_gaps, eval_identity):
        try:
            print(fn())
        except Exception as exc:  # keep going so one failure does not hide the rest
            print(f"{fn.__name__}: ERROR {exc!r}")
    print("=" * 72 + "\n")


if __name__ == "__main__":
    main()
